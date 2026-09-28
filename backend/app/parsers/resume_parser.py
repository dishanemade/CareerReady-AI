"""
Converts an uploaded resume file (PDF or DOCX) into plain text.
This is Step 1 of the pipeline: raw file -> raw text.
"""

import io
from pypdf import PdfReader
from docx import Document


def parse_resume(file_bytes: bytes, filename: str) -> str:
    """
    Extract plain text from a resume file.

    Args:
        file_bytes: raw bytes of the uploaded file
        filename: original filename, used to detect file type

    Returns:
        Extracted plain text as a single string.
    """
    lower_name = filename.lower()

    if lower_name.endswith(".pdf"):
        return _parse_pdf(file_bytes)
    elif lower_name.endswith(".docx"):
        return _parse_docx(file_bytes)
    elif lower_name.endswith(".txt"):
        return file_bytes.decode("utf-8", errors="ignore")
    else:
        raise ValueError(
            f"Unsupported file type for '{filename}'. Please upload a PDF, DOCX, or TXT file."
        )


def _parse_pdf(file_bytes: bytes) -> str:
    reader = PdfReader(io.BytesIO(file_bytes))
    text_parts = []
    for page in reader.pages:
        text = page.extract_text()
        if text:
            text_parts.append(text)
    full_text = "\n".join(text_parts).strip()

    if not full_text:
        raise ValueError(
            "Could not extract any text from this PDF. It may be a scanned "
            "image-based PDF, which requires OCR (not yet supported)."
        )
    return full_text


def _parse_docx(file_bytes: bytes) -> str:
    document = Document(io.BytesIO(file_bytes))
    paragraphs = [p.text for p in document.paragraphs if p.text.strip()]

    # Also pull text out of any tables (some resumes use table layouts)
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    paragraphs.append(cell.text)

    full_text = "\n".join(paragraphs).strip()
    if not full_text:
        raise ValueError("Could not extract any text from this DOCX file.")
    return full_text
